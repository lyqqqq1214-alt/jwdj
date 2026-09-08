from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(r"D:\jwdj\reports\夏季小学期周报第3周_202326010410.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.bold = bold
    r.font.name = '宋体'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader')
    hdr.set(qn('w:val'), 'true')
    trPr.append(hdr)

def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '6')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '808080')
        borders.append(e)
    tbl_pr.append(borders)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

normal = doc.styles['Normal']
normal.font.name = '宋体'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
normal.font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
r = title.add_run('第3周周报')
r.bold = True
r.font.name = '黑体'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
r.font.size = Pt(18)

meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
set_borders(meta)
widths = [Cm(2.3), Cm(5.2), Cm(2.6), Cm(6.4)]
for row in meta.rows:
    for i, c in enumerate(row.cells):
        c.width = widths[i]
        set_cell_margin(c)
for c, text in zip(meta.rows[0].cells, ['姓    名', '202326010410', '参与项目', '基于 AI 的数智化教学分析评价系统']):
    set_cell_text(c, text, bold=text in ['姓    名', '参与项目'], align=WD_ALIGN_PARAGRAPH.CENTER if text in ['姓    名', '参与项目'] else WD_ALIGN_PARAGRAPH.LEFT)
    if text in ['姓    名', '参与项目']:
        set_cell_shading(c, 'EDEDED')
for c, text in zip(meta.rows[1].cells, ['开始时间', '2026-8-31', '结束时间', '2026-9-4']):
    set_cell_text(c, text, bold=text in ['开始时间', '结束时间'], align=WD_ALIGN_PARAGRAPH.CENTER if text in ['开始时间', '结束时间'] else WD_ALIGN_PARAGRAPH.LEFT)
    if text in ['开始时间', '结束时间']:
        set_cell_shading(c, 'EDEDED')

doc.add_paragraph().paragraph_format.space_after = Pt(1)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_borders(table)
col_widths = [Cm(1.2), Cm(10.5), Cm(2.0), Cm(3.0)]
headers = ['序号', '任务内容', '是否完成', '未完成的情况说明']
for i, cell in enumerate(table.rows[0].cells):
    cell.width = col_widths[i]
    set_cell_margin(cell)
    set_cell_shading(cell, 'D9EAF7')
    set_cell_text(cell, headers[i], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(table.rows[0])

tasks = [
    ('1', '在线考试错题沉淀与结果统计：完善考试结果统计逻辑；在学生提交试卷后自动识别错题并写入错题本，保留题干、选项、正确答案、学生答案和知识点快照。', '完成', ''),
    ('2', '错题本与 AI 学习支持：开发错因分析、相似题生成和手动添加错题接口；前端完成错题详情、AI 分析、相似题练习及答错后自动入库流程。', '完成', ''),
    ('3', '本地 AI 服务联调：接入本地 Ollama 模型，调整项目 AI 配置与调用方式；规范 AI 生成题目的题干、选项、答案和解析显示，处理 LaTeX 符号直接暴露的问题。', '完成', ''),
    ('4', '数据库初始化与数据修复：补齐历史数据库缺失字段，修复课程、班级和考试相关接口错误；重新以 UTF-8 方式导入模拟数据和题库数据，核验中文数据与课程学生关联。', '完成', ''),
    ('5', '前后端回归测试与问题修复：验证错题管理、AI 生成、课程选择、班级学生名单等流程；完成 Maven 编译、前端构建和关键接口调用检查。', '完成', ''),
]
for no, task, status, note in tasks:
    cells = table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, [no, task, status, note])):
        cell.width = col_widths[i]
        set_cell_margin(cell)
        set_cell_text(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT)

doc.add_paragraph().paragraph_format.space_after = Pt(2)
summary_table = doc.add_table(rows=1, cols=2)
summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
summary_table.autofit = False
set_borders(summary_table)
summary_table.rows[0].cells[0].width = Cm(2.2)
summary_table.rows[0].cells[1].width = Cm(14.5)
set_cell_margin(summary_table.rows[0].cells[0])
set_cell_margin(summary_table.rows[0].cells[1], top=130, start=150, bottom=130, end=150)
set_cell_shading(summary_table.rows[0].cells[0], 'EDEDED')
set_cell_text(summary_table.rows[0].cells[0], '小    结', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

summary = summary_table.rows[0].cells[1]
summary.text = ''
summary_items = [
    ('在线考试与错题数据闭环：', '本周完善了考试结果统计和错题沉淀逻辑。学生提交试卷后，系统能够记录作答结果并将错误题目自动写入错题本，为后续错因分析和针对性练习提供数据基础。'),
    ('错题本与 AI 功能：', '完成错题本页面与后端接口联调，支持查看题目、正确答案、解析、错因分析和相似题练习；增加手动添加错题功能。相似题练习中若学生再次答错，题目会自动加入错题库，形成持续复习闭环。'),
    ('本地大模型联调与内容优化：', '完成本地 Ollama 服务接入和模型调用验证。针对 AI 输出中出现的 LaTeX 标记和不规范符号，调整生成提示与前端格式化逻辑，使题目、选项、答案和解析以更易读的纯文本形式展示。'),
    ('数据库与测试数据修复：', '排查并修复了数据库字段缺失导致的班级管理、课程加载和考试选课异常。重新以 UTF-8 编码导入模拟数据和题库数据，验证课程、学生、题库中文内容均能正确读取；教师端保留两门可选课程，其中计算机网络班级已关联 68 名学生。'),
    ('总体情况：', '本周围绕在线考试、错题管理、AI 辅助学习和数据稳定性完成了功能实现与联调。前端构建、后端编译及关键接口均已验证通过，系统能够演示从答题、错题沉淀、AI 分析到相似题练习的完整流程。后续将继续完善边界场景和演示数据。'),
]
for idx, (lead, body) in enumerate(summary_items):
    p = summary.add_paragraph() if idx else summary.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(lead)
    r1.bold = True
    r1.font.name = '宋体'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r1.font.size = Pt(10.5)
    r2 = p.add_run(body)
    r2.font.name = '宋体'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r2.font.size = Pt(10.5)

doc.save(OUT)
print(OUT)
